import io
import json
import os
from pathlib import Path

# O pytest-django carrega as configurações antes deste arquivo; por isso o Redis de teste é trocado na fixture
# abaixo (banco 15 por padrão), nunca o banco usado pela aplicação.
TEST_REDIS_URL = os.environ.get("TEST_REDIS_URL", "redis://localhost:6379/15")

import httpx  # noqa: E402
import pytest  # noqa: E402
from rest_framework.test import APIClient  # noqa: E402

from apps.core import redis as r  # noqa: E402

TEXT = (
    "A arquitetura de software organiza um sistema em partes com responsabilidades claras. "
    "Este guia apresenta camadas, filas de mensagens e bancos de dados relacionais com exemplos práticos. "
    "Cada capítulo traz exercícios para equipes que estão começando a documentar decisões técnicas. "
    "Os registros de decisão ajudam a lembrar por que uma escolha foi feita meses depois. "
)


@pytest.fixture(autouse=True)
def _isolation(settings, tmp_path):
    settings.REDIS_URL = TEST_REDIS_URL
    settings.CACHES = {
        "default": {"BACKEND": "django.core.cache.backends.redis.RedisCache", "LOCATION": TEST_REDIS_URL}
    }
    settings.DATA_ROOT = tmp_path / "data"
    settings.DATA_ROOT.mkdir()
    settings.AI = {**settings.AI, "provider": "local", "api_key": "", "model": "", "base_url": ""}
    settings.STORAGES = {**settings.STORAGES, "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"}}
    r.client.cache_clear()
    assert r.client().connection_pool.connection_kwargs.get("db") == int(TEST_REDIS_URL.rsplit("/", 1)[-1])
    r.client().flushdb()
    yield
    r.client().flushdb()
    r.client.cache_clear()


@pytest.fixture
def make_user(db):
    from apps.accounts.models import User

    def factory(email="ana@exemplo.dev", name="Ana Souza", role="member", password="Senha-forte-123"):
        return User.objects.create_user(email=email, password=password, name=name, role=role)

    return factory


@pytest.fixture
def user(make_user):
    return make_user()


@pytest.fixture
def other(make_user):
    return make_user(email="bruno@exemplo.dev", name="Bruno Alves")


@pytest.fixture
def admin(make_user):
    return make_user(email="admin@exemplo.dev", name="Carla Admin", role="admin")


@pytest.fixture
def api(user):
    client = APIClient()
    client.force_login(user)
    return client


@pytest.fixture
def api_for():
    def factory(account):
        client = APIClient()
        client.force_login(account)
        return client

    return factory


def pdf_bytes(pages=3, title="Guia de Arquitetura", author="Equipe Exemplo", text=TEXT) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    c.setTitle(title)
    c.setAuthor(author)
    for page in range(1, pages + 1):
        c.bookmarkPage(f"p{page}")
        c.addOutlineEntry(f"Capítulo {page}", f"p{page}", level=0)
        c.drawString(72, 780, f"Capítulo {page}")
        y = 740
        for line in [text[i : i + 90] for i in range(0, len(text), 90)]:
            c.drawString(72, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buffer.getvalue()


def docx_bytes(title="Manual do Colaborador") -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.title = title
    document.core_properties.author = "RH Exemplo"
    document.add_heading(title, 0)
    document.add_heading("Boas-vindas", 1)
    document.add_paragraph(TEXT)
    document.add_paragraph("Primeiro item", style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def pdf():
    return pdf_bytes


@pytest.fixture
def docx_file():
    return docx_bytes


def upload(client, name, content, **data):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return client.post("/api/v1/documents", {"files": [SimpleUploadedFile(name, content)], **data}, format="multipart")


@pytest.fixture
def send():
    return upload


class FakeProcessor:
    """Simula o processador Java: grava a capa e o HTML na pasta do documento e devolve o JSON do contrato."""

    def __init__(self, settings, status=200, body=None, html=None):
        self.settings = settings
        self.status = status
        self.body = body
        self.html = html
        self.calls = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        self.calls.append(payload)
        if self.status != 200:
            return httpx.Response(self.status, json=self.body or {"error": "internal", "message": "falhou"})
        out = Path(self.settings.DATA_ROOT) / payload["output_dir"]
        out.mkdir(parents=True, exist_ok=True)
        (out / "cover.jpg").write_bytes(b"\xff\xd8\xff\xe0fake-jpeg")
        result = {
            "kind": payload["kind"],
            "page_count": 3,
            "word_count": 120,
            "char_count": 800,
            "title": "Guia de Arquitetura de Software",
            "author": "Equipe Exemplo",
            "subject": None,
            "keywords": None,
            "created_at": "2026-01-10T10:00:00Z",
            "modified_at": None,
            "producer": "ReportLab",
            "creator": None,
            "language": None,
            "outline": [{"title": "Capítulo 1", "page": 1, "level": 1, "anchor": None}],
            "text": "\f".join([TEXT * 3] * 3),
            "text_truncated": False,
            "cover": {"file": "cover.jpg", "width": 480, "height": 679, "color": "#c7b89c", "generated": False},
            "html": None,
            "page_width": 595.0,
            "page_height": 842.0,
            "encrypted": False,
            "warnings": [],
            "elapsed_ms": 42,
        }
        if self.html is not None:
            (out / "reader.html").write_text(self.html, encoding="utf-8")
            result["html"] = {"file": "reader.html", "media": []}
            result["cover"]["generated"] = True
        if self.body:
            result.update(self.body)
        return httpx.Response(200, json=result)

    @property
    def transport(self):
        return httpx.MockTransport(self)


@pytest.fixture
def fake_processor(settings):
    def factory(**kwargs):
        return FakeProcessor(settings, **kwargs)

    return factory
